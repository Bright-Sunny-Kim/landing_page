# -*- coding: utf-8 -*-
"""
감사조서 템플릿(Excel / Word) 전용 파서 및 JSON 생성기
- 대상 경로: data/audit_templates/0_KGAAP_2023
- 출력 경로: data/json_rag/audit_procedure
"""

import os
import sys
import glob
import json
import logging
import re
from typing import Dict, Any, List, Optional
import openpyxl
from openpyxl.worksheet.dimensions import ColumnDimension, RowDimension
import docx

# openpyxl 일부 특수 엑셀 속성(phonetic 등) 호환성 패치
_orig_col_init = ColumnDimension.__init__
def _patched_col_init(self, *args, **kwargs):
    kwargs.pop('phonetic', None)
    _orig_col_init(self, *args, **kwargs)
ColumnDimension.__init__ = _patched_col_init

_orig_row_init = RowDimension.__init__
def _patched_row_init(self, *args, **kwargs):
    kwargs.pop('phonetic', None)
    _orig_row_init(self, *args, **kwargs)
RowDimension.__init__ = _patched_row_init

# 로깅 설정 (Backend Logging Rule 준수)
logging.basicConfig(
    level=logging.INFO,
    format='[%(asctime)s] [%(levelname)s] [%(name)s] %(message)s',
    datefmt='%Y-%m-%d %H:%M:%S'
)
logger = logging.getLogger("AuditTemplateParser")

# 경영진 주장 정규식 및 기본 매핑
ASSERTIONS_MAP = {
    "A": "정확성(Accuracy)",
    "E": "실재성(Existence)",
    "C": "완전성(Completeness)",
    "O": "발생사실(Occurrence)",
    "CO": "기간귀속(Cutoff)",
    "V": "평가(Valuation)",
    "P": "표시 및 공시(Presentation)"
}


def clean_val(val: Any) -> str:
    """셀 값을 문자열로 안전하게 변환하고 공백을 정제합니다."""
    if val is None:
        return ""
    s = str(val).strip()
    return s


def rows_to_markdown_table(rows: List[List[str]]) -> str:
    """2차원 리스트 행들을 Markdown Table 문자열로 변환합니다."""
    if not rows:
        return ""
    
    # 빈 열 제거를 위한 최대 유효 열 계산
    max_cols = max(len(r) for r in rows)
    valid_cols = []
    for c_idx in range(max_cols):
        col_has_val = any(len(r) > c_idx and r[c_idx].strip() for r in rows)
        if col_has_val:
            valid_cols.append(c_idx)
            
    if not valid_cols:
        return ""

    filtered_rows = []
    for r in rows:
        filtered_r = [r[c_idx].replace("\n", " ").replace("|", "\\|") if len(r) > c_idx else "" for c_idx in valid_cols]
        filtered_rows.append(filtered_r)

    if not filtered_rows:
        return ""

    header = filtered_rows[0]
    # 헤더가 완전히 비어있으면 대체 헤더 생성
    if not any(header):
        header = [f"열{i+1}" for i in range(len(valid_cols))]

    md_lines = []
    md_lines.append("| " + " | ".join(header) + " |")
    md_lines.append("| " + " | ".join(["---"] * len(header)) + " |")
    
    for r in filtered_rows[1:]:
        if any(r):  # 완전 빈 행 건너뜀
            md_lines.append("| " + " | ".join(r) + " |")

    return "\n".join(md_lines)


def parse_excel_template(file_path: str, base_dir: str) -> Optional[Dict[str, Any]]:
    """엑셀(.xlsx) 감사조서 파일을 파싱하여 구조화된 딕셔너리로 반환합니다."""
    logger.info("Parsing Excel template: %s", os.path.basename(file_path))
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
    except Exception as e:
        logger.error("Failed to load Excel file: %s, Error: %s", file_path, e, exc_info=True)
        return None

    filename = os.path.basename(file_path)
    rel_path = os.path.relpath(file_path, base_dir)
    path_parts = rel_path.replace("\\", "/").split("/")
    
    section_name = path_parts[0] if len(path_parts) > 1 else "Unknown_Section"
    sub_category = path_parts[1] if len(path_parts) > 2 else ""
    
    # 파일명에서 조서코드 및 제목 추출
    doc_code_match = re.match(r"^([A-Za-z0-9_\-]+)", filename)
    doc_code = doc_code_match.group(1) if doc_code_match else "UNKNOWN"
    
    parsed_sheets = []
    all_instructions = []
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        max_row = ws.max_row or 0
        max_col = ws.max_column or 0
        
        if max_row == 0 or max_col == 0:
            continue

        raw_grid = []
        for r in range(1, max_row + 1):
            row_vals = [clean_val(ws.cell(r, c).value) for c in range(1, max_col + 1)]
            raw_grid.append(row_vals)

        # 1. 메타/지침/절차 구역 분석
        sheet_title = ""
        instructions = []
        assertions_found = set()
        procedures_list = []
        current_category = "일반 절차"
        
        table_rows_buffer = []
        
        for r_idx, row in enumerate(raw_grid):
            row_str = " ".join(v for v in row if v)
            if not row_str:
                continue

            # 제목 탐색 (상단 1~3행)
            if r_idx < 3 and not sheet_title and len(row_str) > 2:
                sheet_title = row_str
                continue

            # 업무요령 / 작성요령 추출
            if "※" in row_str or "업무요령" in row_str or "검토요령" in row_str:
                instructions.append(row_str)
                continue

            # 경영진 주장 감지
            for code in ASSERTIONS_MAP.keys():
                if f"{code}:" in row_str or f"({code})" in row_str:
                    assertions_found.add(code)

            # 섹션 헤더 감지 (예: Part 1, Ⅰ. 일반정보, Ⅱ. 위험평가 등)
            if re.match(r"^(Part\s*\d+|[ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩ]\.|\d+\.\s*[가-힣]+|\[.+?\])", row_str):
                current_category = row_str
                continue

            # 번호가 매겨진 절차 문장 탐색 (예: 1. 또는 1-1. 또는 1) 등)
            step_match = re.match(r"^(\d+[\.\-\)]|\(\d+\))\s*(.+)", row_str)
            if step_match:
                step_no = step_match.group(1)
                proc_text = step_match.group(2)
                
                # 해당 행에 포함된 경영진 주장 추출
                row_assertions = [code for code in ASSERTIONS_MAP.keys() if f" {code} " in f" {row_str} " or f"({code})" in row_str]
                
                procedures_list.append({
                    "step_no": step_no,
                    "category": current_category,
                    "procedure_text": proc_text,
                    "raw_line": row_str,
                    "assertions": row_assertions
                })

            table_rows_buffer.append(row)

        # 표 마크다운 생성
        table_md = rows_to_markdown_table(table_rows_buffer)
        all_instructions.extend(instructions)

        parsed_sheets.append({
            "sheet_name": sheet_name,
            "sheet_title": sheet_title or sheet_name,
            "instructions": "\n".join(instructions),
            "procedures_count": len(procedures_list),
            "procedures": procedures_list,
            "table_markdown": table_md
        })

    result_data = {
        "metadata": {
            "template_id": os.path.splitext(filename)[0],
            "filename": filename,
            "file_type": "xlsx",
            "section_name": section_name,
            "sub_category": sub_category,
            "doc_code": doc_code,
            "year": 2023,
            "total_sheets": len(parsed_sheets)
        },
        "general_instructions": "\n".join(all_instructions),
        "assertions_legend": ASSERTIONS_MAP,
        "sheets": parsed_sheets
    }
    logger.info("Successfully parsed Excel: %s (%d sheets, %d total procedures)",
                filename, len(parsed_sheets), sum(s["procedures_count"] for s in parsed_sheets))
    return result_data


def parse_word_template(file_path: str, base_dir: str) -> Optional[Dict[str, Any]]:
    """워드(.docx) 감사조서 파일을 파싱하여 구조화된 딕셔너리로 반환합니다."""
    logger.info("Parsing Word template: %s", os.path.basename(file_path))
    try:
        doc = docx.Document(file_path)
    except Exception as e:
        logger.error("Failed to load Word file: %s, Error: %s", file_path, e, exc_info=True)
        return None

    filename = os.path.basename(file_path)
    rel_path = os.path.relpath(file_path, base_dir)
    path_parts = rel_path.replace("\\", "/").split("/")
    
    section_name = path_parts[0] if len(path_parts) > 1 else "Unknown_Section"
    sub_category = path_parts[1] if len(path_parts) > 2 else ""

    doc_code_match = re.match(r"^([A-Za-z0-9_\-]+)", filename)
    doc_code = doc_code_match.group(1) if doc_code_match else "UNKNOWN"

    paragraphs_data = []
    tables_data = []
    current_heading = "본문"
    
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if p.style.name.startswith("Heading") or re.match(r"^[ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩ]\.", text):
            current_heading = text
            
        paragraphs_data.append({
            "heading": current_heading,
            "text": text,
            "style": p.style.name
        })

    for t_idx, table in enumerate(doc.tables):
        t_rows = []
        for row in table.rows:
            row_vals = [cell.text.strip() for cell in row.cells]
            t_rows.append(row_vals)
        md_table = rows_to_markdown_table(t_rows)
        tables_data.append({
            "table_index": t_idx + 1,
            "row_count": len(t_rows),
            "table_markdown": md_table
        })

    result_data = {
        "metadata": {
            "template_id": os.path.splitext(filename)[0],
            "filename": filename,
            "file_type": "docx",
            "section_name": section_name,
            "sub_category": sub_category,
            "doc_code": doc_code,
            "year": 2023,
            "total_paragraphs": len(paragraphs_data),
            "total_tables": len(tables_data)
        },
        "paragraphs": paragraphs_data,
        "tables": tables_data
    }
    logger.info("Successfully parsed Word: %s (%d paragraphs, %d tables)",
                filename, len(paragraphs_data), len(tables_data))
    return result_data


def process_all_audit_templates(src_dir: str, out_dir: str) -> Dict[str, int]:
    """전체 감사조서 템플릿 디렉토리를 순회하며 JSON 파일로 변환하여 출력 디렉토리에 저장합니다."""
    logger.info("Starting audit templates batch processing...")
    logger.info("Source Directory: %s", src_dir)
    logger.info("Output Directory: %s", out_dir)

    os.makedirs(out_dir, exist_ok=True)
    
    files = glob.glob(os.path.join(src_dir, "**", "*.*"), recursive=True)
    stats = {"total": 0, "success": 0, "failed": 0, "xlsx": 0, "docx": 0, "skipped": 0}

    for fpath in files:
        ext = os.path.splitext(fpath)[1].lower()
        if ext not in [".xlsx", ".docx"] or os.path.basename(fpath).startswith("~$"):
            stats["skipped"] += 1
            continue

        stats["total"] += 1
        rel_fpath = os.path.relpath(fpath, src_dir)
        rel_dir = os.path.dirname(rel_fpath)
        target_sub_dir = os.path.join(out_dir, rel_dir)
        os.makedirs(target_sub_dir, exist_ok=True)
        
        json_filename = os.path.splitext(os.path.basename(fpath))[0] + ".json"
        target_json_path = os.path.join(target_sub_dir, json_filename)

        parsed_data = None
        if ext == ".xlsx":
            stats["xlsx"] += 1
            parsed_data = parse_excel_template(fpath, src_dir)
        elif ext == ".docx":
            stats["docx"] += 1
            parsed_data = parse_word_template(fpath, src_dir)

        if parsed_data:
            try:
                with open(target_json_path, "w", encoding="utf-8") as jf:
                    json.dump(parsed_data, jf, ensure_ascii=False, indent=2)
                stats["success"] += 1
                logger.debug("Saved JSON to: %s", target_json_path)
            except Exception as se:
                logger.error("Failed to write JSON file: %s, Error: %s", target_json_path, se, exc_info=True)
                stats["failed"] += 1
        else:
            stats["failed"] += 1

    logger.info("==================================================")
    logger.info("Audit Templates Parsing Finished!")
    logger.info("Total Files Evaluated: %d", stats["total"])
    logger.info("  - Excel (.xlsx): %d", stats["xlsx"])
    logger.info("  - Word (.docx): %d", stats["docx"])
    logger.info("Success: %d, Failed: %d, Skipped: %d", stats["success"], stats["failed"], stats["skipped"])
    logger.info("==================================================")
    return stats


if __name__ == "__main__":
    base_proj = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
    source_templates_dir = os.path.join(base_proj, "data", "audit_templates", "0_KGAAP_2023")
    output_rag_dir = os.path.join(base_proj, "data", "json_rag", "audit_procedure")
    
    process_all_audit_templates(source_templates_dir, output_rag_dir)
